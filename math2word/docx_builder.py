"""Structured pages -> editable RTL Word document with native equations."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .measure import line_width_pt
from .omml import latex_to_omml

HEBREW = re.compile(r"[֐-׿יִ-ﭏ]")

# Indents (twips) per nesting level: where the label sits and where the text starts.
LABEL_POS = {0: 0, 1: 0, 2: 567, 3: 1134}
TEXT_POS = {0: 567, 1: 567, 2: 1134, 3: 1701}
# Width estimates are approximate (fonts differ between Word installs): keep a safety margin.
FIT_MARGIN = 0.92


@dataclass
class Options:
    font_latin: str = "Times New Roman"
    font_hebrew: str = "David"
    size_pt: float = 11
    line_spacing: float = 1.5
    page_breaks: bool = True
    margin_cm: float = 1.8
    figure_dpi: int = 300


def _el(tag: str, **attrs) -> OxmlElement:
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


# OOXML is order-sensitive: Word rejects files whose property children are out of schema order.
_ORDER = {
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd "
           "tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi "
           "adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection "
           "textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof "
           "snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText "
           "vertAlign rtl cs em lang eastAsianLayout specVanish oMath",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing "
             "tblInd tblBorders shd tblLayout tblCellMar tblLook",
    "sectPr": "headerReference footerReference footnotePr endnotePr type pgSz pgMar paperSrc pgBorders lnNumType "
              "pgNumType cols formProt vAlign noEndnote titlePg textDirection bidi rtlGutter docGrid printerSettings",
}


def _reorder(parent, kind: str):
    order = {qn("w:" + t): i for i, t in enumerate(_ORDER[kind].split())}
    children = sorted(list(parent), key=lambda c: order.get(c.tag, 999))
    for c in children:
        parent.remove(c)
        parent.append(c)


def _replace(parent, el, kind: str):
    """Insert el into parent (replacing a same-tag child) and restore schema order."""
    for c in parent.findall(el.tag):
        parent.remove(c)
    parent.append(el)
    _reorder(parent, kind)


def _ppr(p):
    return p._p.get_or_add_pPr()


def _set_bidi(p, align: str = "start", level: int = 0, has_label: bool = False, space_before: bool = False,
              keep_next: bool = False):
    pPr = _ppr(p)
    for child in list(pPr):
        if child.tag in (qn("w:bidi"), qn("w:jc"), qn("w:ind"), qn("w:tabs"), qn("w:spacing"), qn("w:keepNext")):
            pPr.remove(child)
    # Schema order: keepNext, ..., tabs, ..., bidi, ..., spacing, ind, ..., jc
    if keep_next:
        pPr.append(_el("w:keepNext"))
    level = max(0, min(3, level))
    text_pos = TEXT_POS[level]
    if has_label:
        tabs = _el("w:tabs")
        tabs.append(_el("w:tab", **{"w:val": "start", "w:pos": text_pos}))
        pPr.append(tabs)
    pPr.append(_el("w:bidi"))
    if space_before:
        pPr.append(_el("w:spacing", **{"w:before": 200}))
    if align != "center" and text_pos:
        ind = {"w:start": text_pos}
        if has_label:
            ind["w:hanging"] = text_pos - LABEL_POS[level]
        pPr.append(_el("w:ind", **ind))
    if align == "center":
        pPr.append(_el("w:jc", **{"w:val": "center"}))
    _reorder(pPr, "pPr")


def _text_run(p, text: str, bold=False, underline=False, size_pt: float | None = None):
    run = p.add_run(text)
    rPr = run._r.get_or_add_rPr()
    if bold:
        rPr.append(_el("w:b"))
        rPr.append(_el("w:bCs"))
    if underline:
        rPr.append(_el("w:u", **{"w:val": "single"}))
    if size_pt:
        rPr.append(_el("w:sz", **{"w:val": int(size_pt * 2)}))
        rPr.append(_el("w:szCs", **{"w:val": int(size_pt * 2)}))
    # Every text run in a Hebrew paragraph is a right-to-left run: keeps punctuation,
    # parentheses and spaces on the correct side of neighbouring equations.
    rPr.append(_el("w:rtl"))
    _reorder(rPr, "rPr")
    return run


def _size_math(omath, size_pt: float):
    """Set an explicit font size on every math run (w:rPr goes after m:rPr, before m:t)."""
    M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
    for r in omath.iter(M + "r"):
        for old in r.findall(qn("w:rPr")):
            r.remove(old)
        rPr = _el("w:rPr")
        rPr.append(_el("w:sz", **{"w:val": int(size_pt * 2)}))
        rPr.append(_el("w:szCs", **{"w:val": int(size_pt * 2)}))
        mrpr = r.find(M + "rPr")
        r.insert(1 if mrpr is not None else 0, rPr)


class _MathPool:
    """Collects all LaTeX first, converts in one batch, then hands out OMML elements."""

    def __init__(self, pages: list[dict]):
        formulas: list[str] = []

        def walk_lines(lines):
            for line in lines:
                for seg in line:
                    if seg["type"] == "math" and seg["value"].strip():
                        formulas.append(seg["value"])

        for page in pages:
            for b in page.get("blocks", []):
                walk_lines(b.get("lines", []))
                for par in b.get("paragraphs", []):
                    walk_lines(par.get("lines", []))
                for row in b.get("rows", []):
                    walk_lines(row)
        uniq = list(dict.fromkeys(formulas))
        self._map = dict(zip(uniq, latex_to_omml(uniq)))

    def get(self, latex: str):
        import copy
        e = self._map.get(latex)
        return copy.deepcopy(e) if e is not None else None


class Builder:
    def __init__(self, options: Options | None = None):
        self.o = options or Options()
        self.doc = Document()
        self._setup()

    # ---------- document setup ----------
    def _setup(self):
        o = self.o
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Cm(o.margin_cm))
        _replace(sec._sectPr, _el("w:bidi"), "sectPr")

        styles = self.doc.styles.element
        rpr_default = styles.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
        for child in list(rpr_default):
            rpr_default.remove(child)
        rpr_default.append(_el("w:rFonts", **{"w:ascii": o.font_latin, "w:hAnsi": o.font_latin,
                                               "w:cs": o.font_hebrew, "w:eastAsia": o.font_latin}))
        rpr_default.append(_el("w:sz", **{"w:val": int(o.size_pt * 2)}))
        rpr_default.append(_el("w:szCs", **{"w:val": int(o.size_pt * 2)}))
        rpr_default.append(_el("w:lang", **{"w:val": "en-US", "w:bidi": "he-IL"}))

        zoom = self.doc.settings.element.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")

        normal = self.doc.styles["Normal"]
        pf = normal.paragraph_format
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        pf.line_spacing = o.line_spacing
        _replace(normal.element.get_or_add_pPr(), _el("w:bidi"), "pPr")


    # ---------- content ----------
    def _fill_lines(self, p, lines: list, bold_all=False, size_pt=None):
        for i, line in enumerate(lines):
            if i:
                p.add_run().add_break(WD_BREAK.LINE)
            for seg in line:
                if seg["type"] == "math":
                    omml = self.math.get(seg["value"])
                    if omml is not None:
                        if size_pt:
                            _size_math(omml, size_pt)
                        p._p.append(omml)
                    else:  # conversion failed: keep the LaTeX source visible for manual fixing
                        run = p.add_run(seg["value"])
                        run.italic = True
                        _reorder(run._r.get_or_add_rPr(), "rPr")
                else:
                    _text_run(p, seg["value"], bold=seg.get("bold") or bold_all,
                              underline=seg.get("underline"), size_pt=size_pt)

    # ---------- line-length preservation ----------
    def _needed_pt(self, par: dict, size_pt: float) -> float:
        """Width a paragraph needs so that none of its source lines wraps."""
        level = max(0, min(3, par.get("level", 0)))
        widest = max((line_width_pt(l, self.o.font_hebrew, size_pt) for l in par.get("lines", [])), default=0)
        return (0 if par.get("align") == "center" else TEXT_POS[level] / 20) + widest

    def _fit_size(self, par: dict, avail_pt: float, base: float) -> float | None:
        """Font size that keeps every line on one line (None = base size fits)."""
        need = self._needed_pt(par, base)
        limit = avail_pt * FIT_MARGIN
        if need <= limit:
            return None
        size = base * limit / need
        return max(base - 2.5, int(size * 2) / 2)

    def _paragraph(self, container, par: dict, heading=False, avail_pt: float | None = None):
        p = container.add_paragraph()
        label = (par.get("label") or "").strip()
        _set_bidi(p, par.get("align", "start"), par.get("level", 0), bool(label),
                  par.get("space_before", False) or heading, keep_next=heading)
        base = self.o.size_pt + 1 if heading else self.o.size_pt
        size = self._fit_size(par, avail_pt, base) if avail_pt else None
        size = size or (base if heading else None)
        if label:
            _text_run(p, label, bold=True, size_pt=size)
            p.add_run().add_tab()
        self._fill_lines(p, par.get("lines", []), bold_all=heading, size_pt=size)
        return p

    def _spacer(self):
        """Tiny paragraph (Word needs one between consecutive tables)."""
        p = self.doc.add_paragraph()
        _set_bidi(p)
        pPr = _ppr(p)
        rPr = _el("w:rPr")
        rPr.append(_el("w:sz", **{"w:val": 4}))
        rPr.append(_el("w:szCs", **{"w:val": 4}))
        pPr.append(rPr)
        _replace(pPr, _el("w:spacing", **{"w:line": 240, "w:lineRule": "auto"}), "pPr")

    def _picture(self, container, page_obj, fig: dict, max_width_pt: float, align_center=True, crop=None):
        img, (w_pt, h_pt) = crop or page_obj.crop_figure(fig["bbox"], dpi=self.o.figure_dpi)
        scale = min(1.0, max_width_pt / w_pt) if w_pt else 1.0
        buf = io.BytesIO()
        img.save(buf, format="PNG", dpi=(self.o.figure_dpi, self.o.figure_dpi))
        buf.seek(0)
        p = container.add_paragraph()
        _set_bidi(p, "center" if align_center else "start")
        run = p.add_run()
        shape = run.add_picture(buf, width=Pt(w_pt * scale))
        docpr = shape._inline.docPr
        docpr.set("descr", fig.get("description", ""))
        return p, w_pt * scale

    def _table_props(self, table, widths_twips: list[int], borders: bool):
        tbl = table._tbl
        tblPr = tbl.tblPr
        b = _el("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b.append(_el(f"w:{side}", **({"w:val": "single", "w:sz": 4, "w:space": 0, "w:color": "000000"}
                                          if borders else {"w:val": "nil"})))
        for el in (_el("w:bidiVisual"), _el("w:tblW", **{"w:w": sum(widths_twips), "w:type": "dxa"}),
                   _el("w:tblLayout", **{"w:type": "fixed"}), b):
            _replace(tblPr, el, "tblPr")
        grid = tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths_twips):
            gc.set(qn("w:w"), str(w))
        for row in table.rows:
            for cell, w in zip(row.cells, widths_twips):
                cell.width = Pt(w / 20)

    def _clear_cell(self, cell):
        for p in list(cell.paragraphs):
            p._p.getparent().remove(p._p)

    def add_page(self, page_json: dict, page_obj, first: bool):
        text_width_pt = (21 - 2 * self.o.margin_cm) / 2.54 * 72
        blocks = page_json.get("blocks", [])
        if not blocks:
            return False
        if not first and self.o.page_breaks:
            self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            _set_bidi(self.doc.paragraphs[-1])
        figures = {f["id"]: f for f in page_json.get("figures", [])}
        for b in blocks:
            kind = b["kind"]
            if kind in ("paragraph", "heading"):
                self._paragraph(self.doc, b, heading=(kind == "heading"), avail_pt=text_width_pt)
            elif kind == "display_math":
                self._paragraph(self.doc, {**b, "align": "center", "label": b.get("label", "")}, avail_pt=text_width_pt)
            elif kind == "figure" and b["figure_id"] in figures:
                self._picture(self.doc, page_obj, figures[b["figure_id"]], text_width_pt)
            elif kind == "figure_row" and b["figure_id"] in figures:
                fig = figures[b["figure_id"]]
                table = self.doc.add_table(rows=1, cols=2)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                # bidiVisual: the first cell is rendered on the RIGHT.
                fig_first = b.get("figure_side") == "right"
                fig_cell, txt_cell = (table.cell(0, 0), table.cell(0, 1)) if fig_first else (table.cell(0, 1), table.cell(0, 0))
                self._clear_cell(fig_cell)
                self._clear_cell(txt_cell)
                crop = page_obj.crop_figure(fig["bbox"], dpi=self.o.figure_dpi)
                native_w = min(crop[1][0], text_width_pt * 0.45)
                # Give the text what it needs to keep its lines; the figure may shrink to 70 %.
                cell_pad, gap = 11, 14
                need = max((self._needed_pt(par, self.o.size_pt) for par in b.get("paragraphs", [])), default=0) / FIT_MARGIN
                fig_w = native_w
                if text_width_pt - native_w - gap - cell_pad < need:
                    fig_w = max(0.7 * native_w, text_width_pt - gap - cell_pad - need)
                _, fig_w = self._picture(fig_cell, page_obj, fig, fig_w, crop=crop)
                fig_tw = int((fig_w + gap) * 20)
                total_tw = int(text_width_pt * 20)
                widths = [fig_tw, total_tw - fig_tw] if fig_first else [total_tw - fig_tw, fig_tw]
                self._table_props(table, widths, borders=False)
                avail = (total_tw - fig_tw) / 20 - cell_pad
                for par in b.get("paragraphs", []):
                    self._paragraph(txt_cell, par, avail_pt=avail)
                if not txt_cell.paragraphs:
                    txt_cell.add_paragraph()
                self._spacer()
            elif kind == "table" and b.get("rows"):
                rows = b["rows"]
                ncols = max(len(r) for r in rows)
                table = self.doc.add_table(rows=len(rows), cols=ncols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                col_tw = int(min(text_width_pt, ncols * 90) * 20 / ncols)
                self._table_props(table, [col_tw] * ncols, borders=True)
                for r, row in enumerate(rows):
                    for c, cell_line in enumerate(row):
                        cell = table.cell(r, c)
                        p = cell.paragraphs[0]
                        _set_bidi(p, "center")
                        self._fill_lines(p, [cell_line])
                self._spacer()
        return True

    def build(self, pages_json: list[dict], page_objs: list, out_path, include_types=("questions", "instructions",
                                                                                          "formula_sheet", "other")):
        self.math = _MathPool(pages_json)
        first = True
        for pj, po in zip(pages_json, page_objs):
            if pj.get("page_type") not in include_types:
                continue
            if self.add_page(pj, po, first):
                first = False
        # python-docx starts with one empty paragraph? (no) – ensure body is valid and save.
        self.doc.save(out_path)
        return out_path
